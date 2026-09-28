"""Read the radar's report contract without exposing collector credentials."""
from __future__ import annotations

import math
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook

MAX_UPLOAD = 10 * 1024 * 1024
FIELDS = {
    'Produto': 'title', 'Categoria': 'category', 'Tipo': 'type', 'ID': 'id',
    'Posição bestseller': 'rank', 'Preço observado (R$)': 'price',
    'Avaliações': 'reviews', 'Nota média': 'rating', 'Demanda (0-100)': 'demand',
    'Cobertura demanda (%)': 'demandCoverage', 'Potencial de marca (0-100)': 'brand',
    'Cobertura marca (%)': 'brandCoverage', 'Prioridade combinada': 'combined',
    'Evidências de demanda': 'demandEvidence', 'Hipóteses de diferenciação': 'brandEvidence',
    'Contrapontos': 'counterpoints', 'Link': 'url',
}
NUMERIC = {'rank', 'price', 'reviews', 'rating', 'demand', 'demandCoverage', 'brand', 'brandCoverage', 'combined'}
SCORES = {'demand', 'brand', 'combined', 'demandCoverage', 'brandCoverage'}

CATEGORY_NAMES = {
    "MLB109197": "Moldes para Artesanatos",
    "MLB1166": "Pelúcias",
    "MLB180337": "Slackline",
    "MLB270072": "Mesas e Cadeiras",
    "MLB270271": "Anilinas",
    "MLB271735": "Elásticos",
    "MLB31767": "Placas",
    "MLB381266": "Espelhos de Mosaico",
    "MLB3900": "Tênis",
    "MLB39492": "Brinquedos de Pegadinhas",
    "MLB432171": "Massas de Biscuit",
    "MLB432184": "Placas de Cortiça",
    "MLB434742": "Adesivos para Frasco e Garrafa",
    "MLB436797": "Enfeites e Objetos para Pintar",
    "MLB455695": "Correntes",
    "MLB455730": "Kits Pontilhismo",
    "MLB455774": "Massinhas de Modelar",
}



def read_report(path: Path) -> dict:
    try:
        with ZipFile(path) as archive:
            if len(archive.infolist()) > 300 or sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Planilha excede o limite de conteúdo expandido.')
        source = path.open('rb')
        try:
            book = load_workbook(source, read_only=True, data_only=True)
        except Exception:
            source.close()
            raise
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError('Arquivo XLSX inválido.') from exc
    try:
        if not {'Oportunidades', 'Investigar', 'Execução'}.issubset(book.sheetnames):
            raise ValueError('Use um relatório do Radar com as abas Oportunidades, Investigar e Execução.')
        metadata = book['Execução'].iter_rows(values_only=True)
        info = dict(zip(next(metadata, ()), next(metadata, ())))
        run_id = str(info.get('ID') or '')
        if not run_id or len(run_id) > 128:
            raise ValueError('Identificador da coleta ausente ou inválido.')
        try:
            date = datetime.fromisoformat(str(info['Início UTC']).replace('Z', '+00:00'))
            if date.tzinfo is None:
                date = date.replace(tzinfo=timezone.utc)
        except (ValueError, KeyError) as exc:
            raise ValueError('Data da coleta inválida.') from exc
        products = []
        seen = set()
        for sheet, priority in [('Oportunidades', True), ('Investigar', False)]:
            rows = book[sheet].iter_rows(values_only=True)
            headers = next(rows, ())
            for row in rows:
                if not any(x is not None for x in row):
                    continue
                if not {'Produto', 'ID', 'Tipo', 'Categoria', 'Demanda (0-100)', 'Potencial de marca (0-100)'}.issubset(headers):
                    raise ValueError('Colunas de produtos incompatíveis com o Radar.')
                p = {target: dict(zip(headers, row)).get(source) for source, target in FIELDS.items()}
                for key, value in p.items():
                    if key in NUMERIC:
                        if value is not None:
                            try:
                                value = float(value)
                            except (TypeError, ValueError) as exc:
                                raise ValueError('Valor numérico inválido.') from exc
                            if not math.isfinite(value) or value < 0 or (key in SCORES and value > 100) or (key == 'rating' and value > 5):
                                raise ValueError('Valor numérico fora do intervalo permitido.')
                        p[key] = value
                    else:
                        p[key] = str(value or '')[:12000]
                if not p['id'] or not p['title']:
                    raise ValueError('Produto sem ID ou nome.')
                key = p['type'] + ':' + p['id']
                if key in seen:
                    continue
                seen.add(key)
                p['key'] = key
                p['priority'] = priority
                if not p['url'].startswith('https://'):
                    p['url'] = ''
                if p.get('category') in CATEGORY_NAMES:
                    p['category'] = CATEGORY_NAMES[p['category']]
                products.append(p)
                if len(products) > 10000:
                    raise ValueError('Limite de 10.000 produtos por relatório.')
        if not products:
            raise ValueError('O relatório não contém produtos.')
        return {'id': run_id, 'date': date.astimezone(timezone.utc).isoformat(),
                'status': str(info.get('Status') or 'unknown'), 'mode': str(info.get('Modo') or 'importado'),
                'products': products, 'filename': path.name}
    finally:
        book.close()
        source.close()


def reports_in(directory: Path) -> tuple[list[dict], int]:
    runs = {}
    skipped = 0
    for path in sorted(directory.glob('*.xlsx')):
        try:
            report = read_report(path)
            runs.setdefault(report['id'], report)
        except Exception:
            skipped += 1
    return sorted(runs.values(), key=lambda r: (r['date'], r['id'])), skipped
