from pathlib import Path
from typing import Union
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from src.arbitrage.models import ArbitrageReport

def export_arbitrage_excel(report: ArbitrageReport, output_path: Union[str, Path] = "arbitragem_alibaba_ml.xlsx") -> Path:
    """Exports an ArbitrageReport into a beautifully formatted Excel workbook."""
    path = Path(output_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Oportunidades de Arbitragem"

    # Header fonts and fills
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="176449", end_color="176449", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    headers = [
        "Estratégia",
        "Produto (China)",
        "Fonte",
        "FOB (US$)",
        "FOB (R$)",
        "MOQ",
        "Custo Brasil (R$)",
        "Preço Médio ML (R$)",
        "Taxa ML (R$)",
        "Lucro Líquido (R$)",
        "Margem Líquida (%)",
        "ROI (%)",
        "Concorrentes ML",
        "% Full no ML",
        "Fornecedor",
        "Diagnóstico",
        "Link China",
        "Link Concorrente ML",
    ]

    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Strategy color fills for data rows
    strategy_colors = {
        "OCEANO_AZUL": PatternFill(start_color="E8F8F0", end_color="E8F8F0", fill_type="solid"),   # soft green
        "ALTA_MARGEM": PatternFill(start_color="E6F0FA", end_color="E6F0FA", fill_type="solid"),   # soft blue
        "MARCA_PROPRIA": PatternFill(start_color="F5EEF8", end_color="F5EEF8", fill_type="solid"), # soft purple
        "SATURADO": PatternFill(start_color="FDF2E9", end_color="FDF2E9", fill_type="solid"),      # soft orange
    }

    strategy_font_colors = {
        "OCEANO_AZUL": Font(color="0E6251", bold=True),
        "ALTA_MARGEM": Font(color="1B4F72", bold=True),
        "MARCA_PROPRIA": Font(color="512E5F", bold=True),
        "SATURADO": Font(color="784212", bold=True),
    }

    for row_idx, item in enumerate(report.items, start=2):
        ws.row_dimensions[row_idx].height = 22
        row_data = [
            item.strategy_label,
            item.china_title,
            item.source,
            item.china_fob_usd,
            item.china_fob_brl,
            item.moq,
            item.landed_cost_unit_brl,
            item.reference_sale_price_brl,
            item.ml_fee_brl,
            item.net_profit_brl,
            f"{item.net_margin_pct:.1f}%",
            f"{item.roi_pct:.1f}%",
            item.ml_competitors_count,
            f"{item.ml_full_share_pct:.0f}%",
            item.supplier_name,
            item.summary,
            item.china_url,
            item.ml_top_competitor_url,
        ]
        ws.append(row_data)

        # Style strategy column specially
        strat_cell = ws.cell(row=row_idx, column=1)
        if item.strategy in strategy_colors:
            strat_cell.fill = strategy_colors[item.strategy]
            strat_cell.font = strategy_font_colors[item.strategy]
        strat_cell.alignment = center_align

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.border = thin_border
            if col_idx in (4, 5, 7, 8, 9, 10):
                c.alignment = right_align
            elif col_idx in (1, 3, 6, 11, 12, 13, 14):
                c.alignment = center_align
            else:
                c.alignment = left_align

    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(12, min(50, max_len + 3))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path
